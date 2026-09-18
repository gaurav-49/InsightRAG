"""End-to-end check against the Docker Compose stack (design doc §8, end-to-end row).

Upload through to answered query, including the NO_ANSWER path and the dead-letter path, with
the real API, the real worker, Postgres/pgvector and Redis.

    docker compose up -d --build
    python scripts/e2e.py

Needs httpx, reportlab and python-docx, plus the worker package importable (to compute the
expected chunk count with the worker's own chunker): run it with worker/.venv/bin/python.
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import time
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))
sys.path.insert(0, str(REPO / "worker"))

from insightrag_client import API, Client  # noqa: E402
from insightrag_worker.chunker import ChunkingConfig, chunk_document  # noqa: E402
from insightrag_worker.extract import extract  # noqa: E402
from insightrag_worker.tokenizer import RegexTokenizer  # noqa: E402

CHUNK_SIZE = int(os.environ.get("CHUNK_SIZE_TOKENS", "128"))
CHUNK_OVERLAP = int(os.environ.get("CHUNK_OVERLAP_TOKENS", "0"))

passed = 0


def check(cond: bool, what: str, detail=None) -> None:
    global passed
    if not cond:
        print(f"FAIL  {what}" + (f"\n      {json.dumps(detail, default=str)[:800]}" if detail is not None else ""))
        sys.exit(1)
    passed += 1
    print(f"ok    {what}")


def make_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in text.splitlines():
        line = line.replace("#", "").strip()
        while line:
            c.drawString(50, y, line[:100])
            line = line[100:]
            y -= 14
            if y < 60:
                c.showPage()
                y = 800
    c.save()
    return buf.getvalue()


def make_docx(text: str) -> bytes:
    import docx

    d = docx.Document()
    for block in text.split("\n\n"):
        block = block.strip()
        if block.startswith("#"):
            d.add_heading(block.lstrip("# "), level=min(3, len(block) - len(block.lstrip("#"))))
        elif block:
            d.add_paragraph(block)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def compose(*args: str) -> str:
    return subprocess.run(["docker", "compose", *args], cwd=REPO, check=True, capture_output=True, text=True).stdout


def main() -> int:
    admin = Client(API, Client.dev_token(API, ["admin", "query"]))
    health = admin.wait_healthy()
    check(health["status"] == "UP", "stack healthy", health)

    # Journey A — every supported format.
    spec = (REPO / "eval" / "corpus" / "product-spec-halcyon-arm-x2.md").read_text()
    pdf = make_pdf(spec)
    r = admin.upload(f"arm-x2-spec-{int(time.time())}.pdf", pdf, "e2e")
    check(r.status_code in (200, 202), "PDF upload accepted without waiting for processing", r.text)
    pdf_id = r.json()["documentId"]
    if r.status_code == 202:
        check(r.json()["status"] == "QUEUED", "upload returns QUEUED immediately", r.json())

    r = admin.upload("again.pdf", pdf, "e2e")
    check(r.status_code == 200 and r.json()["duplicate"] and r.json()["documentId"] == pdf_id,
          "identical bytes short-circuit to the existing document (200)", r.json())

    docx_id = admin.upload("security-policy.docx", make_docx((REPO / "eval/corpus/information-security-policy.md").read_text()),
                           "e2e").json()["documentId"]
    corpus_ids = []
    for path in sorted((REPO / "eval" / "corpus").iterdir()):
        resp = admin.upload(path.name, path.read_bytes(), "eval-corpus")
        check(resp.status_code in (200, 202), f"upload {path.name}", resp.text)
        corpus_ids.append(resp.json()["documentId"])

    states = admin.wait_terminal([pdf_id, docx_id, *corpus_ids])
    expected = len(chunk_document(extract(pdf, "application/pdf"), RegexTokenizer(), ChunkingConfig(CHUNK_SIZE, CHUNK_OVERLAP)))
    check(states[pdf_id]["status"] == "INDEXED" and states[pdf_id]["chunkCount"] == expected,
          f"PDF reaches INDEXED with the correct chunk count ({expected})", states[pdf_id])
    check(states[docx_id]["status"] == "INDEXED", "DOCX reaches INDEXED", states[docx_id])
    check(all(states[d]["status"] == "INDEXED" for d in corpus_ids), "Markdown and TXT corpus indexed",
          {d: states[d]["status"] for d in corpus_ids})

    # Journey B — grounded, cited answer; then the cache tiers.
    user = Client(API, Client.dev_token(API, ["query"]))
    q = "What is the maximum payload of the ARM-X2?"
    first = user.query(q).json()
    check(first["status"] == "ANSWERED" and first["citations"], "answerable question is answered with citations", first)
    check("7 kilograms" in first["answer"], "answer is grounded in the retrieved passage", first["answer"])
    check(any("arm-x2" in c["document"].lower() for c in first["citations"]), "citation names the source document",
          first["citations"])
    second = user.query(q.upper()).json()
    check(second["cacheTier"] == "L1" and second["cached"], "repeat question served from L1 without an LLM call", second)
    check(second["latencyMs"] < 200, f"cache hit latency {second['latencyMs']} ms < 200 ms")
    third = user.query("What is the maximum payload of the ARM-X2, then?").json()
    check(third["cacheTier"] == "L2", "near-duplicate served from the semantic tier", third)

    none = user.query("What is the recipe for sourdough bread?").json()
    check(none["status"] == "NO_ANSWER" and none["answer"] is None and none["citations"] == [],
          "unanswerable question is refused, not fabricated", none)

    filtered = user.query("What is the maximum payload of the ARM-X2?", source="no-such-source").json()
    check(filtered["status"] == "NO_ANSWER", "metadata filter applied in retrieval", filtered)

    # FR-15 streaming
    with user.http.stream("POST", "/api/v1/query/stream", json={"question": "How long is the standard warranty on the ARM-X2?"}) as s:
        body = "".join(s.iter_text())
    check("event:meta" in body and "event:token" in body and "event:done" in body, "answer streams as server-sent events",
          body[:300])

    # Failure paths.
    bad = admin.upload("corrupt.pdf", b"%PDF-1.4\n this is not a real pdf " + str(time.time()).encode())
    bad_state = admin.wait_terminal([bad.json()["documentId"]])[bad.json()["documentId"]]
    check(bad_state["status"] == "FAILED" and bad_state["failureReason"], "malformed document fails with a readable reason",
          bad_state)

    dlq_before = int(compose("exec", "-T", "redis", "redis-cli", "XLEN", "insightrag:ingest:dlq").strip() or 0)
    compose("exec", "-T", "redis", "redis-cli", "XADD", "insightrag:ingest", "*", "documentId", "not-a-uuid", "storageKey", "x")
    deadline = time.time() + 30
    dlq_after = dlq_before
    while time.time() < deadline and dlq_after == dlq_before:
        time.sleep(1)
        dlq_after = int(compose("exec", "-T", "redis", "redis-cli", "XLEN", "insightrag:ingest:dlq").strip() or 0)
    check(dlq_after == dlq_before + 1, "poison message moved to the dead-letter stream without blocking the queue")

    # FR-12 delete + invalidation
    r = admin.http.delete(f"/api/v1/documents/{pdf_id}")
    check(r.status_code == 204, "document deleted")
    after = user.query(q).json()
    check(after["cacheTier"] == "MISS", "corpus change invalidated the cached answer", after)

    metrics = admin.http.get("/api/v1/metrics").text
    for name in ("insightrag_queries_total", "insightrag_ingest_queue_lag", "insightrag_retrieval_top_score",
                 "insightrag_llm_calls_total"):
        check(name in metrics, f"metric {name} exposed")

    # Leave only the labelled corpus behind, so a following eval run scores exactly eval/corpus.
    for extra in (docx_id, bad.json()["documentId"]):
        admin.http.delete(f"/api/v1/documents/{extra}")

    print(f"\nall {passed} end-to-end checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
