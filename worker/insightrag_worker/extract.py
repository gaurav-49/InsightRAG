"""Text extraction for PDF, DOCX, TXT and Markdown (FR-01).

Extraction failures raise :class:`ExtractionError` with a message written for the corpus
administrator; the processor stores it verbatim as ``failure_reason`` (§5.7, "Malformed or
unparseable document").
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import List, Optional, Tuple


class ExtractionError(Exception):
    """Permanent: retrying the same bytes cannot succeed."""


@dataclass
class ExtractedDocument:
    text: str
    # Character offset at which each page starts (index 0 => page 1). Empty when the format
    # has no pages.
    page_starts: List[int] = field(default_factory=list)
    # (offset, heading) pairs in document order.
    headings: List[Tuple[int, str]] = field(default_factory=list)

    def page_at(self, offset: int) -> Optional[int]:
        if not self.page_starts:
            return None
        page = 1
        for i, start in enumerate(self.page_starts):
            if start <= offset:
                page = i + 1
            else:
                break
        return page

    def heading_at(self, offset: int) -> Optional[str]:
        current = None
        for start, heading in self.headings:
            if start <= offset:
                current = heading
            else:
                break
        return current


_MD_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t#]*$", re.MULTILINE)
_NUMBERED_HEADING = re.compile(r"^(\d+(\.\d+)*\.?|[A-Z]\.)\s+[A-Z][^.!?]{1,80}$")

MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
}


def extract(data: bytes, mime_type: str) -> ExtractedDocument:
    if mime_type == "application/pdf":
        doc = _pdf(data)
    elif mime_type == MIME_BY_EXTENSION[".docx"]:
        doc = _docx(data)
    elif mime_type == "text/markdown":
        doc = _markdown(_decode(data))
    elif mime_type == "text/plain":
        doc = ExtractedDocument(text=_normalise_newlines(_decode(data)))
    else:
        raise ExtractionError(f"Unsupported document type {mime_type!r}.")
    if not doc.text.strip():
        raise ExtractionError(
            "No extractable text was found. Scanned or image-only documents are not supported in version 1."
        )
    return doc


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError(f"Text file is not valid UTF-8 (invalid byte at position {exc.start}).") from exc


def _normalise_newlines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _markdown(text: str) -> ExtractedDocument:
    text = _normalise_newlines(text)
    headings = [(m.start(), m.group(2).strip()) for m in _MD_HEADING.finditer(text)]
    return ExtractedDocument(text=text, headings=headings)


def _looks_like_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 90 or s.endswith((".", ",", ";", ":")):
        return False
    if _NUMBERED_HEADING.match(s):
        return True
    words = s.split()
    return 1 <= len(words) <= 8 and (s.isupper() and any(c.isalpha() for c in s))


def _pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("PDF is encrypted; upload an unencrypted copy.")
        pages = [(p.extract_text() or "") for p in reader.pages]
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise ExtractionError(f"PDF could not be parsed ({type(exc).__name__}). The file may be corrupt.") from exc

    parts: List[str] = []
    page_starts: List[int] = []
    headings: List[Tuple[int, str]] = []
    offset = 0
    for page_text in pages:
        page_text = _normalise_newlines(page_text).strip()
        page_starts.append(offset)
        line_offset = offset
        for line in page_text.split("\n"):
            if _looks_like_heading(line):
                headings.append((line_offset, line.strip()))
            line_offset += len(line) + 1
        parts.append(page_text)
        offset += len(page_text) + 2  # the "\n\n" page separator
    return ExtractedDocument(text="\n\n".join(parts), page_starts=page_starts, headings=headings)


def _docx(data: bytes) -> ExtractedDocument:
    import zipfile

    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("DOCX could not be opened; the file is not a valid Word document.") from exc

    parts: List[str] = []
    headings: List[Tuple[int, str]] = []
    offset = 0

    def add(text: str) -> None:
        nonlocal offset
        parts.append(text)
        offset += len(text) + 2

    from docx.table import Table

    # iter_inner_content() yields paragraphs and tables in body order, so table text lands
    # under the heading it appears beneath.
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    add(" | ".join(cells))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = ((block.style.name if block.style is not None else "") or "").lower()
        if style.startswith("heading") or style == "title":
            headings.append((offset, text))
        add(text)
    return ExtractedDocument(text="\n\n".join(parts), headings=headings)
