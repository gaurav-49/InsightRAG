import io

import pytest

from insightrag_worker.extract import MIME_BY_EXTENSION, ExtractionError, extract


def make_pdf(pages):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for lines in pages:
        y = 740
        for line in lines:
            c.drawString(72, y, line)
            y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx():
    import docx

    d = docx.Document()
    d.add_heading("Expense Policy", level=1)
    d.add_paragraph("Meals are reimbursed up to 75 dollars per day.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Limit"
    t.cell(1, 0).text, t.cell(1, 1).text = "Hotel", "250"
    d.add_heading("Approvals", level=2)
    d.add_paragraph("Managers approve claims within five business days.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_pages_and_offsets():
    data = make_pdf([["1. Introduction", "The policy applies to all staff."],
                     ["2. Notice Period", "Employees give sixty days notice."]])
    doc = extract(data, "application/pdf")
    assert len(doc.page_starts) == 2
    idx = doc.text.index("sixty days")
    assert doc.page_at(idx) == 2
    assert doc.page_at(doc.text.index("all staff")) == 1
    assert doc.heading_at(idx) == "2. Notice Period"


def test_docx_headings_and_tables_in_order():
    doc = extract(make_docx(), MIME_BY_EXTENSION[".docx"])
    assert "Hotel | 250" in doc.text
    assert doc.text.index("Hotel | 250") < doc.text.index("Approvals")
    assert doc.heading_at(doc.text.index("Hotel")) == "Expense Policy"
    assert doc.heading_at(doc.text.index("five business days")) == "Approvals"


def test_markdown_headings():
    doc = extract(b"# Title\n\ntext\n\n## Sub ##\n\nmore", "text/markdown")
    assert [h for _, h in doc.headings] == ["Title", "Sub"]


def test_txt_crlf_normalised_and_bom_stripped():
    doc = extract(b"\xef\xbb\xbfline one\r\nline two", "text/plain")
    assert doc.text == "line one\nline two"


@pytest.mark.parametrize("data,mime,fragment", [
    (b"\xff\xfe\x00bad", "text/plain", "not valid UTF-8"),
    (b"%PDF-1.4 garbage", "application/pdf", "PDF could not be parsed"),
    (b"not a zip", MIME_BY_EXTENSION[".docx"], "not a valid Word document"),
    (b"   ", "text/plain", "No extractable text"),
    (b"x", "image/png", "Unsupported"),
])
def test_failures_are_human_readable(data, mime, fragment):
    with pytest.raises(ExtractionError, match=fragment):
        extract(data, mime)


def test_image_only_pdf_reports_no_text():
    with pytest.raises(ExtractionError, match="No extractable text"):
        extract(make_pdf([[]]), "application/pdf")
